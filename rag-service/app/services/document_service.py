import io
from pathlib import Path

import docx
import pypdf
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

SUPPORTED_EXTENSIONS = {".txt": "text", ".md": "text", ".pdf": "pdf", ".docx": "docx"}


class DocumentParseError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            f"不支持的文档类型: {ext or '无扩展名'}，支持 {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    try:
        if ext in (".txt", ".md"):
            return data.decode("utf-8", errors="replace")
        if ext == ".pdf":
            reader = pypdf.PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(page.strip() for page in pages if page.strip())
        if ext == ".docx":
            document = docx.Document(io.BytesIO(data))
            paragraphs = [
                para.text.strip()
                for para in document.paragraphs
                if para.text and para.text.strip()
            ]
            return "\n\n".join(paragraphs)
    except Exception as exc:
        raise DocumentParseError(f"文档解析失败: {exc}") from exc
    raise DocumentParseError("文档内容为空")


def split_documents(
    filename: str,
    data: bytes,
    document_id: str,
    chunk_size: int,
    chunk_overlap: int,
) -> list[Document]:
    text = extract_text(filename, data)
    if not text.strip():
        raise DocumentParseError("文档内容为空")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", "。", "；", "，", " ", ""],
    )
    chunks = splitter.split_text(text)
    return [
        Document(
            page_content=chunk,
            metadata={"document_id": document_id, "source": Path(filename).name},
        )
        for chunk in chunks
    ]